import io
import json
import os
import random
import secrets
from copy import copy
from datetime import datetime, timedelta, timezone
from functools import wraps

from flask import (
    Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
)
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook, load_workbook
from sqlalchemy import UniqueConstraint, func
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash
from werkzeug.utils import secure_filename

db = SQLAlchemy()


class Course(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(160), nullable=False)
    code = db.Column(db.String(30), unique=True, nullable=False, index=True)
    duration_minutes = db.Column(db.Integer, nullable=False, default=15)
    questions_per_student = db.Column(db.Integer, nullable=False, default=20)
    pass_mark = db.Column(db.Float, nullable=False, default=50.0)
    is_active = db.Column(db.Boolean, nullable=False, default=False)
    created_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    questions = db.relationship("Question", backref="course", cascade="all, delete-orphan")
    attempts = db.relationship("Attempt", backref="course", cascade="all, delete-orphan")


class Question(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    course_id = db.Column(db.Integer, db.ForeignKey("course.id"), nullable=False, index=True)
    text = db.Column(db.Text, nullable=False)
    options_json = db.Column(db.Text, nullable=False)
    correct_answer = db.Column(db.Text, nullable=False)
    mark = db.Column(db.Float, nullable=False, default=1.0)

    @property
    def options(self):
        return json.loads(self.options_json)


class Attempt(db.Model):
    __table_args__ = (
        UniqueConstraint("course_id", "matric_number", name="uq_course_matric"),
        UniqueConstraint("course_id", "exam_number", name="uq_course_exam"),
    )
    id = db.Column(db.Integer, primary_key=True)
    public_id = db.Column(db.String(64), unique=True, nullable=False, index=True)
    course_id = db.Column(db.Integer, db.ForeignKey("course.id"), nullable=False, index=True)
    full_name = db.Column(db.String(160), nullable=False)
    matric_number = db.Column(db.String(80), nullable=False)
    phone_number = db.Column(db.String(30), nullable=False)
    email = db.Column(db.String(160), nullable=False)
    exam_number = db.Column(db.String(80), nullable=False)
    question_order_json = db.Column(db.Text, nullable=False)
    option_order_json = db.Column(db.Text, nullable=False)
    answers_json = db.Column(db.Text, nullable=False, default="{}")
    score = db.Column(db.Float)
    total_marks = db.Column(db.Float)
    percentage = db.Column(db.Float)
    status = db.Column(db.String(20), nullable=False, default="in_progress")
    started_at = db.Column(db.DateTime(timezone=True), nullable=False)
    deadline_at = db.Column(db.DateTime(timezone=True), nullable=False)
    submitted_at = db.Column(db.DateTime(timezone=True))


def normalize_database_url(value):
    value = value or "sqlite:///cbt_local.db"
    if value.startswith("postgres://"):
        value = value.replace("postgres://", "postgresql+psycopg://", 1)
    elif value.startswith("postgresql://") and "+psycopg" not in value:
        value = value.replace("postgresql://", "postgresql+psycopg://", 1)
    return value


def create_app(test_config=None):
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.getenv("SECRET_KEY", secrets.token_hex(32)),
        SQLALCHEMY_DATABASE_URI=normalize_database_url(os.getenv("DATABASE_URL")),
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        MAX_CONTENT_LENGTH=5 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.getenv("FLASK_ENV") == "production",
    )
    if test_config:
        app.config.update(test_config)
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)
    db.init_app(app)

    with app.app_context():
        db.create_all()

    def csrf_token():
        if "csrf_token" not in session:
            session["csrf_token"] = secrets.token_urlsafe(32)
        return session["csrf_token"]

    app.jinja_env.globals["csrf_token"] = csrf_token

    @app.before_request
    def validate_csrf():
        if request.method == "POST":
            supplied = request.form.get("csrf_token", "")
            expected = session.get("csrf_token", "")
            if not expected or not secrets.compare_digest(supplied, expected):
                abort(400, "Invalid or expired form token. Refresh the page and try again.")

    def admin_required(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not session.get("admin_authenticated"):
                flash("Please sign in to access the administrator area.", "warning")
                return redirect(url_for("admin_login", next=request.path))
            return view(*args, **kwargs)
        return wrapped

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.post("/access-test")
    def access_test():
        code = request.form.get("course_code", "").strip().upper()
        course = Course.query.filter(func.upper(Course.code) == code, Course.is_active.is_(True)).first()
        if not course:
            flash("No active test was found for that course code.", "danger")
            return redirect(url_for("index"))
        if len(course.questions) < course.questions_per_student:
            flash("This test is not ready. Please contact the administrator.", "danger")
            return redirect(url_for("index"))
        return redirect(url_for("student_details", code=course.code))

    @app.route("/test/<code>/details", methods=["GET", "POST"])
    def student_details(code):
        course = Course.query.filter(func.upper(Course.code) == code.upper(), Course.is_active.is_(True)).first_or_404()
        if request.method == "POST":
            fields = {
                "full_name": request.form.get("full_name", "").strip(),
                "matric_number": request.form.get("matric_number", "").strip().upper(),
                "phone_number": request.form.get("phone_number", "").strip(),
                "email": request.form.get("email", "").strip().lower(),
                "exam_number": request.form.get("exam_number", "").strip().upper(),
            }
            if not all(fields.values()):
                flash("All student information fields are required.", "danger")
                return render_template("student_details.html", course=course, values=fields)
            existing = Attempt.query.filter(
                Attempt.course_id == course.id,
                db.or_(Attempt.matric_number == fields["matric_number"], Attempt.exam_number == fields["exam_number"]),
            ).first()
            if existing:
                flash("An attempt already exists for this matriculation number or examination number.", "danger")
                return render_template("student_details.html", course=course, values=fields)

            selected = random.SystemRandom().sample(course.questions, course.questions_per_student)
            option_order = {}
            for question in selected:
                shuffled = list(question.options)
                random.SystemRandom().shuffle(shuffled)
                option_order[str(question.id)] = shuffled
            now = datetime.now(timezone.utc)
            attempt = Attempt(
                public_id=secrets.token_urlsafe(24), course_id=course.id,
                question_order_json=json.dumps([q.id for q in selected]),
                option_order_json=json.dumps(option_order), started_at=now,
                deadline_at=now + timedelta(minutes=course.duration_minutes), **fields
            )
            db.session.add(attempt)
            db.session.commit()
            session["attempt_public_id"] = attempt.public_id
            return redirect(url_for("take_test", public_id=attempt.public_id))
        return render_template("student_details.html", course=course, values={})

    def get_owned_attempt(public_id):
        attempt = Attempt.query.filter_by(public_id=public_id).first_or_404()
        if session.get("attempt_public_id") != public_id:
            abort(403)
        return attempt

    @app.get("/attempt/<public_id>")
    def take_test(public_id):
        attempt = get_owned_attempt(public_id)
        if attempt.status == "completed":
            return redirect(url_for("result", public_id=public_id))
        question_ids = json.loads(attempt.question_order_json)
        questions_by_id = {q.id: q for q in Question.query.filter(Question.id.in_(question_ids)).all()}
        questions = [questions_by_id[qid] for qid in question_ids]
        options = json.loads(attempt.option_order_json)
        deadline = attempt.deadline_at
        if deadline.tzinfo is None:
            deadline = deadline.replace(tzinfo=timezone.utc)
        seconds_remaining = max(0, int((deadline - datetime.now(timezone.utc)).total_seconds()))
        return render_template("take_test.html", attempt=attempt, questions=questions,
                               options=options, seconds_remaining=seconds_remaining)

    @app.post("/attempt/<public_id>/submit")
    def submit_test(public_id):
        attempt = get_owned_attempt(public_id)
        if attempt.status == "completed":
            return redirect(url_for("result", public_id=public_id))
        question_ids = json.loads(attempt.question_order_json)
        questions = Question.query.filter(Question.id.in_(question_ids)).all()
        answers = {str(qid): request.form.get(f"question_{qid}", "") for qid in question_ids}
        score = sum(q.mark for q in questions if answers.get(str(q.id)) == q.correct_answer)
        total = sum(q.mark for q in questions)
        attempt.answers_json = json.dumps(answers)
        attempt.score = round(score, 2)
        attempt.total_marks = round(total, 2)
        attempt.percentage = round((score / total * 100) if total else 0, 2)
        attempt.status = "completed"
        attempt.submitted_at = datetime.now(timezone.utc)
        db.session.commit()
        return redirect(url_for("result", public_id=public_id))

    @app.get("/attempt/<public_id>/result")
    def result(public_id):
        attempt = get_owned_attempt(public_id)
        if attempt.status != "completed":
            return redirect(url_for("take_test", public_id=public_id))
        return render_template("result.html", attempt=attempt)

    @app.route("/admin/login", methods=["GET", "POST"])
    def admin_login():
        if request.method == "POST":
            username = request.form.get("username", "")
            password = request.form.get("password", "")
            configured_user = os.getenv("ADMIN_USERNAME", "admin")
            password_hash = os.getenv("ADMIN_PASSWORD_HASH")
            plain_password = os.getenv("ADMIN_PASSWORD")
            valid_password = (password_hash and check_password_hash(password_hash, password)) or (
                plain_password and secrets.compare_digest(plain_password, password)
            )
            if secrets.compare_digest(username, configured_user) and valid_password:
                session.clear()
                session["admin_authenticated"] = True
                csrf_token()
                return redirect(request.args.get("next") or url_for("admin_dashboard"))
            flash("Incorrect administrator username or password.", "danger")
        return render_template("admin_login.html")

    @app.post("/admin/logout")
    @admin_required
    def admin_logout():
        session.clear()
        return redirect(url_for("index"))

    @app.get("/admin")
    @admin_required
    def admin_dashboard():
        courses = Course.query.order_by(Course.created_at.desc()).all()
        totals = {
            "courses": Course.query.count(),
            "active": Course.query.filter_by(is_active=True).count(),
            "attempts": Attempt.query.filter_by(status="completed").count(),
        }
        return render_template("admin_dashboard.html", courses=courses, totals=totals)

    @app.post("/admin/courses/create")
    @admin_required
    def create_course():
        title = request.form.get("title", "").strip()
        code = request.form.get("code", "").strip().upper()
        if not title or not code:
            flash("Test title and course code are required.", "danger")
        elif Course.query.filter(func.upper(Course.code) == code).first():
            flash("That course code already exists.", "danger")
        else:
            course = Course(title=title, code=code, duration_minutes=15,
                            questions_per_student=20, pass_mark=50)
            db.session.add(course)
            db.session.commit()
            flash("Course created. Upload at least 20 valid questions before activation.", "success")
        return redirect(url_for("admin_dashboard"))

    @app.get("/admin/courses/<int:course_id>")
    @admin_required
    def manage_course(course_id):
        course = Course.query.get_or_404(course_id)
        search = request.args.get("q", "").strip()
        query = Attempt.query.filter_by(course_id=course.id)
        if search:
            pattern = f"%{search}%"
            query = query.filter(db.or_(Attempt.full_name.ilike(pattern), Attempt.matric_number.ilike(pattern),
                                        Attempt.exam_number.ilike(pattern), Attempt.email.ilike(pattern)))
        attempts = query.order_by(Attempt.started_at.desc()).limit(100).all()
        return render_template("manage_course.html", course=course, attempts=attempts, search=search)

    @app.post("/admin/courses/<int:course_id>/upload")
    @admin_required
    def upload_questions(course_id):
        course = Course.query.get_or_404(course_id)
        uploaded = request.files.get("question_file")
        if not uploaded or not uploaded.filename:
            flash("Choose an Excel file to upload.", "danger")
            return redirect(url_for("manage_course", course_id=course.id))
        filename = secure_filename(uploaded.filename)
        if not filename.lower().endswith(".xlsx"):
            flash("Only .xlsx files are accepted.", "danger")
            return redirect(url_for("manage_course", course_id=course.id))
        try:
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
            sheet = workbook.active
            headers = [str(c.value or "").strip().lower() for c in sheet[1]]
            required = ["question", "option a", "option b", "option c", "option d", "correct answer", "mark"]
            if headers[:7] != required:
                raise ValueError("The first seven columns must match the sample template.")
            imported = []
            seen = set()
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                values = list(row[:7])
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                question = str(values[0] or "").strip()
                options = [str(v or "").strip() for v in values[1:5]]
                answer = str(values[5] or "").strip()
                if not question or any(not option for option in options):
                    raise ValueError(f"Row {row_number}: question and four options are required.")
                if len(set(options)) != 4:
                    raise ValueError(f"Row {row_number}: all four options must be different.")
                if answer.upper() in {"A", "B", "C", "D"}:
                    answer = options[ord(answer.upper()) - ord("A")]
                if answer not in options:
                    raise ValueError(f"Row {row_number}: correct answer must be A-D or exactly match an option.")
                mark = float(values[6] if values[6] is not None else 1)
                if mark <= 0:
                    raise ValueError(f"Row {row_number}: mark must be greater than zero.")
                key = question.casefold()
                if key in seen:
                    raise ValueError(f"Row {row_number}: duplicate question in the uploaded file.")
                seen.add(key)
                imported.append(Question(course_id=course.id, text=question,
                                         options_json=json.dumps(options), correct_answer=answer, mark=mark))
            if len(imported) < course.questions_per_student:
                raise ValueError(f"Upload at least {course.questions_per_student} valid questions.")
            Question.query.filter_by(course_id=course.id).delete()
            db.session.add_all(imported)
            course.is_active = False
            db.session.commit()
            flash(f"{len(imported)} questions imported successfully. Preview and activate the test.", "success")
        except Exception as exc:
            db.session.rollback()
            flash(f"Upload failed: {exc}", "danger")
        return redirect(url_for("manage_course", course_id=course.id))

    @app.post("/admin/courses/<int:course_id>/toggle")
    @admin_required
    def toggle_course(course_id):
        course = Course.query.get_or_404(course_id)
        if len(course.questions) < course.questions_per_student:
            flash(f"At least {course.questions_per_student} questions are required before activation.", "danger")
        else:
            course.is_active = not course.is_active
            db.session.commit()
            flash(f"Test {'activated' if course.is_active else 'deactivated'}.", "success")
        return redirect(url_for("manage_course", course_id=course.id))

    @app.get("/admin/courses/<int:course_id>/results.xlsx")
    @admin_required
    def export_results(course_id):
        course = Course.query.get_or_404(course_id)
        attempts = Attempt.query.filter_by(course_id=course.id, status="completed").order_by(Attempt.submitted_at).all()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"
        sheet.append(["Full Name", "Matriculation Number", "Phone Number", "Email Address",
                      "Exam Number", "Course Code", "Score", "Total Marks", "Percentage",
                      "Status", "Started At", "Submitted At"])
        for attempt in attempts:
            sheet.append([attempt.full_name, attempt.matric_number, attempt.phone_number, attempt.email,
                          attempt.exam_number, course.code, attempt.score, attempt.total_marks,
                          attempt.percentage, "Pass" if attempt.percentage >= course.pass_mark else "Fail",
                          attempt.started_at.replace(tzinfo=None),
                          attempt.submitted_at.replace(tzinfo=None) if attempt.submitted_at else None])
        for cell in sheet[1]:
            new_font = copy(cell.font)
            new_font.bold = True
            new_font.color = "FFFFFF"
            cell.font = new_font
            new_fill = copy(cell.fill)
            new_fill.fill_type = "solid"
            new_fill.fgColor.rgb = "173F35"
            cell.fill = new_fill
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(36, max(12, max(len(str(c.value or "")) for c in column) + 2))
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return send_file(buffer, as_attachment=True,
                         download_name=f"{course.code}_CBT_results.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.get("/health")
    def health():
        return {"status": "ok"}

    return app


app = create_app()


if __name__ == "__main__":
    app.run(debug=True)
